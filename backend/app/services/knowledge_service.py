"""知识文档上传、解析、切块与向量摄取服务。by AI.Coding"""

from __future__ import annotations

import hashlib
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Protocol

from docx import Document
from pypdf import PdfReader
from sqlmodel import Session, col, delete, select

from app.core.config import settings
from app.core.errors import ConflictError, ErrorCode, NotFoundError, ValidationError
from app.core.workspace import WorkspaceContext, WorkspacePolicy
from app.models.knowledge import (
    DocumentIngestionJob,
    IngestionJobStatus,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeDocumentStatus,
)
from app.models.user import User
from app.providers.embedding import EmbeddingProvider

SUPPORTED_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
}
# 浏览器在 Windows 上可能把文档标记为空 MIME 或 octet-stream；这些值仍可由扩展名安全识别。by AI.Coding
UPLOAD_MIME_ALIASES = {
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".markdown": {"text/markdown", "text/plain", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
}
CHUNK_SIZE_CHARS = 2800
CHUNK_OVERLAP_CHARS = 400


class KnowledgeUpload(Protocol):
    """上传对象所需的最小接口，便于路由和单测复用。by AI.Coding"""

    filename: str | None
    content_type: str | None

    async def read(self, size: int = -1) -> bytes:
        """读取上传内容。by AI.Coding"""
        ...


@dataclass(frozen=True)
class ParsedSegment:
    """解析后的文本与来源定位信息。by AI.Coding"""

    text: str
    source_meta: dict[str, object]


def validate_upload_metadata(
    filename: str | None,
    content_type: str | None,
    size_bytes: int,
) -> tuple[str, str]:
    """校验文件名、MIME 与大小并返回安全扩展名和标准 MIME。by AI.Coding"""
    if not filename:
        raise ValidationError(ErrorCode.VALIDATION_ERROR)
    safe_name = Path(filename).name
    if safe_name != filename or safe_name in {".", ".."}:
        raise ValidationError(ErrorCode.VALIDATION_ERROR)
    extension = Path(safe_name).suffix.lower()
    expected_mime = SUPPORTED_MIME_TYPES.get(extension)
    if expected_mime is None or size_bytes <= 0 or size_bytes > settings.KNOWLEDGE_MAX_FILE_BYTES:
        raise ValidationError(ErrorCode.VALIDATION_ERROR)
    supplied_mime = (content_type or "").split(";", 1)[0].strip().lower()
    # 不能要求浏览器上报 MIME 必须精确一致，否则合法的 Markdown/DOCX 会被 422 拒绝。by AI.Coding
    allowed_mimes = UPLOAD_MIME_ALIASES[extension]
    if supplied_mime and supplied_mime not in allowed_mimes:
        raise ValidationError(ErrorCode.VALIDATION_ERROR)
    return extension, expected_mime


def chunk_text(
    text: str,
    *,
    chunk_size: int = CHUNK_SIZE_CHARS,
    overlap: int = CHUNK_OVERLAP_CHARS,
) -> list[str]:
    """按近似 token 窗口切分文本并保留有限重叠。by AI.Coding"""
    normalized = "\n".join(line.strip() for line in text.splitlines()).strip()
    if not normalized:
        return []
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("切块参数无效")
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = end - overlap
    return chunks


def parse_document(content: bytes, extension: str) -> list[ParsedSegment]:
    """按受支持格式解析文本并附带页码或段落来源。by AI.Coding"""
    if extension in {".txt", ".md", ".markdown"}:
        text = content.decode("utf-8", errors="replace")
        return [ParsedSegment(text=text, source_meta={"kind": "document"})]
    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))
        segments = [
            ParsedSegment(
                text=page.extract_text() or "",
                source_meta={"kind": "page", "page": index + 1},
            )
            for index, page in enumerate(reader.pages)
        ]
        return segments
    if extension == ".docx":
        document = Document(BytesIO(content))
        return [
            ParsedSegment(
                text=paragraph.text,
                source_meta={"kind": "paragraph", "paragraph": index + 1},
            )
            for index, paragraph in enumerate(document.paragraphs)
        ]
    raise ValueError("不支持的文档格式")


class KnowledgeService:
    """管理 Workspace 知识文档及其持久化摄取任务。by AI.Coding"""

    def __init__(self, session: Session) -> None:
        """保存请求级数据库会话。by AI.Coding"""
        self.session = session

    async def create_document(
        self,
        context: WorkspaceContext,
        actor: User,
        upload: KnowledgeUpload,
    ) -> tuple[KnowledgeDocument, DocumentIngestionJob]:
        """校验并持久化文档元数据，上传请求只创建任务不执行解析。by AI.Coding"""
        WorkspacePolicy.require_manager(context)
        # 保留原始文件名交给校验器，不能先规范化后再校验，否则会掩盖路径穿越。by AI.Coding
        filename = upload.filename or ""
        content = await upload.read(settings.KNOWLEDGE_MAX_FILE_BYTES + 1)
        extension, mime_type = validate_upload_metadata(
            filename, upload.content_type, len(content)
        )
        safe_filename = Path(filename).name
        document_id = uuid.uuid4()
        storage_key = f"{context.workspace_id}/{document_id}{extension}"
        storage_path = self._storage_path(storage_key)
        storage_path.parent.mkdir(parents=True, exist_ok=True)
        storage_path.write_bytes(content)
        document = KnowledgeDocument(
            id=document_id,
            workspace_id=context.workspace_id,
            display_name=safe_filename,
            storage_key=storage_key,
            mime_type=mime_type,
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            created_by_id=actor.id,
        )
        job = DocumentIngestionJob(
            workspace_id=context.workspace_id,
            document_id=document_id,
        )
        try:
            self.session.add(document)
            # 文档与摄取任务没有 ORM relationship，必须先 flush 文档，避免任务先插入触发外键 500。by AI.Coding
            self.session.flush()
            self.session.add(job)
            self.session.commit()
        except Exception:
            self.session.rollback()
            storage_path.unlink(missing_ok=True)
            raise
        self.session.refresh(document)
        self.session.refresh(job)
        return document, job

    def get_document(
        self, context: WorkspaceContext, document_id: uuid.UUID
    ) -> KnowledgeDocument:
        """按 Workspace 读取未删除文档，跨租户统一隐藏。by AI.Coding"""
        document = self.session.exec(
            select(KnowledgeDocument).where(
                KnowledgeDocument.id == document_id,
                KnowledgeDocument.workspace_id == context.workspace_id,
                KnowledgeDocument.status != KnowledgeDocumentStatus.DELETED,
            )
        ).first()
        if document is None:
            raise NotFoundError(ErrorCode.WORKSPACE_RESOURCE_NOT_FOUND)
        return document

    def retry_document(
        self, context: WorkspaceContext, document_id: uuid.UUID
    ) -> DocumentIngestionJob:
        """为失败文档创建幂等摄取任务并恢复 PROCESSING 状态。by AI.Coding"""
        WorkspacePolicy.require_manager(context)
        document = self.get_document(context, document_id)
        if document.status is not KnowledgeDocumentStatus.FAILED:
            raise ConflictError(ErrorCode.PROVIDER_UNAVAILABLE)
        job = DocumentIngestionJob(
            workspace_id=context.workspace_id,
            document_id=document.id,
        )
        document.status = KnowledgeDocumentStatus.PROCESSING
        document.error_code = None
        self.session.add(document)
        self.session.add(job)
        self.session.commit()
        self.session.refresh(job)
        return job

    def list_documents(
        self, context: WorkspaceContext
    ) -> list[KnowledgeDocument]:
        """返回当前 Workspace 尚未删除的知识文档。by AI.Coding"""
        WorkspacePolicy.require_manager(context)
        return list(
            self.session.exec(
                select(KnowledgeDocument)
                .where(
                    KnowledgeDocument.workspace_id == context.workspace_id,
                    KnowledgeDocument.status != KnowledgeDocumentStatus.DELETED,
                )
                .order_by(col(KnowledgeDocument.created_at).desc())
            ).all()
        )

    def delete_document(
        self, context: WorkspaceContext, document_id: uuid.UUID
    ) -> None:
        """软删除文档并移除其可检索切块，保留历史任务记录。by AI.Coding"""
        WorkspacePolicy.require_manager(context)
        document = self.get_document(context, document_id)
        document.status = KnowledgeDocumentStatus.DELETED
        document.deleted_at = document.deleted_at or datetime.now(UTC)
        self.session.exec(
            delete(KnowledgeChunk).where(
                col(KnowledgeChunk.document_id) == document.id,
                col(KnowledgeChunk.workspace_id) == context.workspace_id,
            )
        )
        self.session.add(document)
        self.session.commit()
        self._storage_path(document.storage_key).unlink(missing_ok=True)

    def mark_failed(
        self,
        context: WorkspaceContext,
        document_id: uuid.UUID,
        *,
        error_code: str,
    ) -> None:
        """将文档和任务标记为可重试的 FAILED 状态。by AI.Coding"""
        document = self.get_document(context, document_id)
        document.status = KnowledgeDocumentStatus.FAILED
        document.error_code = error_code
        job = self.session.exec(
            select(DocumentIngestionJob)
            .where(
                DocumentIngestionJob.document_id == document_id,
                DocumentIngestionJob.workspace_id == context.workspace_id,
            )
            .order_by(col(DocumentIngestionJob.created_at).desc())
        ).first()
        if job is not None:
            job.status = IngestionJobStatus.FAILED
            job.error_code = error_code
        self.session.commit()

    def ingest_document(
        self,
        context: WorkspaceContext,
        document_id: uuid.UUID,
        embedding_provider: EmbeddingProvider,
    ) -> KnowledgeDocument:
        """解析文档、切块、批量生成向量并写入 pgvector。by AI.Coding"""
        document = self.get_document(context, document_id)
        job = self.session.exec(
            select(DocumentIngestionJob)
            .where(
                DocumentIngestionJob.document_id == document.id,
                DocumentIngestionJob.workspace_id == context.workspace_id,
                col(DocumentIngestionJob.status).in_(
                    [IngestionJobStatus.PENDING, IngestionJobStatus.RUNNING]
                ),
            )
            .order_by(col(DocumentIngestionJob.created_at).desc())
        ).first()
        if job is None:
            raise NotFoundError(ErrorCode.WORKSPACE_RESOURCE_NOT_FOUND)
        job.status = IngestionJobStatus.RUNNING
        job.attempts += 1
        self.session.add(job)
        self.session.commit()
        try:
            content = self._storage_path(document.storage_key).read_bytes()
            segments = parse_document(content, Path(document.display_name).suffix.lower())
            chunks = [
                (chunk, segment.source_meta)
                for segment in segments
                for chunk in chunk_text(segment.text)
            ]
            if not chunks:
                raise ValueError("文档没有可检索文本")
            vectors = self._embed_chunks(embedding_provider, [item[0] for item in chunks])
            self.session.exec(
                delete(KnowledgeChunk).where(
                    col(KnowledgeChunk.document_id) == document.id,
                    col(KnowledgeChunk.workspace_id) == context.workspace_id,
                )
            )
            for index, ((chunk, source_meta), vector) in enumerate(zip(chunks, vectors, strict=True)):
                self.session.add(
                    KnowledgeChunk(
                        workspace_id=context.workspace_id,
                        document_id=document.id,
                        chunk_index=index,
                        content=chunk,
                        source_meta=source_meta,
                        embedding=vector,
                    )
                )
            document.status = KnowledgeDocumentStatus.READY
            document.error_code = None
            job.status = IngestionJobStatus.READY
            self.session.add(document)
            self.session.add(job)
            self.session.commit()
            self.session.refresh(document)
            return document
        except Exception:
            self.session.rollback()
            self.mark_failed(context, document_id, error_code="DOCUMENT_PROCESSING_FAILED")
            raise

    @staticmethod
    def _embed_chunks(
        embedding_provider: EmbeddingProvider, texts: list[str]
    ) -> list[list[float]]:
        """同步 worker 中调用异步 Embedding provider 并校验数量。by AI.Coding"""
        import asyncio

        vectors = asyncio.run(embedding_provider.embed(texts))
        if len(vectors) != len(texts):
            raise ValueError("Embedding 返回数量不一致")
        return vectors

    @staticmethod
    def _storage_path(storage_key: str) -> Path:
        """将相对存储键解析到知识库根目录并阻断路径穿越。by AI.Coding"""
        root = settings.KNOWLEDGE_STORAGE_DIR.resolve()
        path = (root / storage_key).resolve()
        if root not in path.parents:
            raise ValueError("非法存储路径")
        return path
